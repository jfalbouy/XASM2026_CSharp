from __future__ import annotations

from datetime import date
from pathlib import Path
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Racines deduites de l'emplacement du script : le generateur suit le depot dans lequel
# il se trouve, au lieu de dependre d'une arborescence machine (anciennement C:\Codex\...).
ROOT = Path(__file__).resolve().parent.parent
FINAL_ROOT = Path(os.environ.get("XASM_FINAL_ROOT", str(ROOT)))
DOC_DIR = ROOT / "Documentation"
FINAL_DOC_DIR = FINAL_ROOT / "Documentation"

# Ressources externes absentes du depot : a fournir via variable d'environnement.
# Les valeurs par defaut ne pointent volontairement sur aucune arborescence
# personnelle : ce sont des emplacements relatifs au depot, a surcharger par
# XASM_INSTRUCTION_TABLE / XASM140_INFO_TXT selon la machine.
INSTRUCTION_TABLE = Path(os.environ.get(
    "XASM_INSTRUCTION_TABLE",
    str(ROOT / "_ressources_externes" / "README - PC-E500 Instruction Table.md")))
XASM140_INFO_TXT = Path(os.environ.get(
    "XASM140_INFO_TXT",
    str(ROOT / "_ressources_externes" / "XASM140 - Information.DOC.txt")))


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def fenced(text: str) -> str:
    return "```text\n" + text.strip("\n") + "\n```"


def build_xasm2026_3_markdown() -> str:
    instruction = read(INSTRUCTION_TABLE).strip()
    today = date.today().strftime("%d/%m/%Y")
    return f"""# Documentation XASM2026-4 pour Sharp PC-E500S

Version du document : {today}

## 1. Objet du document

Ce document décrit `xasm2026-4`, port C# natif de l'assembleur XASM pour CPU Sharp ESR-L / SC62015. Il reprend la documentation pratique établie pour `xasm2026-1`, ajoute les choix de portage effectués pour `xasm2026-4`, décrit les options de sortie modernes et inclut en annexe le contenu complet du fichier `README - PC-E500 Instruction Table.md`.

`xasm2026-4` a été construit pour conserver le comportement observable de XASM 1.40 et du moteur C `xasm2026-2`, tout en remplaçant progressivement le coeur assembleur par du code C# maintenable.

| Élément | Valeur |
|---|---|
| Projet | `xasm2026-4` |
| Langage | C# / .NET 8 |
| CPU cible | Sharp ESR-L / SC62015 |
| Machines visées | Sharp PC-E500 / PC-E500S et proches compatibles |
| Référence historique | XASM 1.40 |
| Référence de comparaison | `xasm2026-2`, XASM 1.40 sous DOSBox et sorties historiques |
| Exécutable principal | `bin\\xasm2026-4.exe` |

## 2. Arborescence du projet

| Chemin | Rôle |
|---|---|
| `src/` | Projet C# natif. |
| `src/Assembly/NativeAssembler.cs` | Lecture source, macros, conditions, labels, encodage SC62015. |
| `src/Outputs/` | Générateurs `.obj`, `.lst`, `.hex`, `.s19`, `.map`, `.d`, `.uu`, `.txt`. |
| `bin/` | Copie pratique de l'exécutable et des fichiers runtime. |
| `Exemples/` | Sources et sorties de validation : `SAMPLES`, `REGISTER`, `VOGUE`, `TRDOS`, `UUCODE`. |
| `Reference/C/` | Sources C historiques utilisées comme référence de portage. |
| `Reference/CSharpWrapper/` | Trace de l'étape `xasm2026-2`, wrapper autour du moteur C. |
| `tests/` | Fichiers de couverture et tests de non-régression. |
| `Documentation/` | Documentation Markdown et Word. |

## 3. Compilation du projet

Depuis la racine du projet :

```powershell
dotnet build .\\src\\Xasm2026.Native.csproj -c Release
```

Le build génère :

```text
src\\bin\\Release\\net8.0\\xasm2026-4.exe
```

Une copie pratique est maintenue ici :

```text
bin\\xasm2026-4.exe
```

## 4. Ligne de commande

Forme générale :

```powershell
xasm2026-4.exe sourcefile[.ext] [options]
```

Exemple complet :

```powershell
xasm2026-4.exe coverage_all.asm -O coverage_all.obj -L coverage_all.lst -E -S -TZ -C -W -H -I coverage_all.hex -M coverage_all.s19 -P coverage_all.map -D coverage_all.d -B coverage_all.uu -X coverage_all.txt -V -R
```

| Option | Effet |
|---|---|
| `-L[filename]` | Génère le listing assembleur. |
| `-E` | Génère le rapport d'erreur `.err`; en cas d'erreur fatale, le `.lst` reçoit aussi le rapport si `-L` est actif. |
| `-O[filename]` | Génère le fichier objet. |
| `-S` | Active la liste de symboles dans le listing. |
| `-T[type]` | Type objet : `Z`, `B`, `H`, `F` ou défaut avec en-tête. |
| `-C` | Affiche le nombre de lignes traitées. |
| `-W` | Active les avertissements compatibles avec l'interface historique. |
| `-H` | Désactive le hash dans l'esprit de l'option historique. |
| `-I[filename]` | Génère un fichier Intel HEX. |
| `-M[filename]` | Génère un Motorola S-Record. |
| `-P[filename]` | Génère la MAP symboles/sections. |
| `-D[filename]` | Génère les dépendances, notamment les fichiers inclus. |
| `-B[filename]` | Génère un BASIC auto-décodeur UUENCODE. |
| `-X[filename]` | Génère un dump texte façon HxD. |
| `-V` | Active les diagnostics détaillés. |
| `-R` | Affiche le rapport de taille par section. |

## 5. Syntaxe source

Chaque ligne contient au plus une instruction ou directive. La forme usuelle est :

```asm
LABEL:  MNEMONIC  operand1,operand2   ; commentaire
        MNEMONIC  operand
        ; commentaire
        END
```

Règles principales :

- `END` est obligatoire dans le fichier principal et dans les fichiers inclus historiques.
- Les labels sont insensibles à la casse.
- Les commentaires commencent par `;`.
- Les opérandes peuvent contenir des constantes, labels, expressions, caractères, chaînes et compteur de position.
- Les constantes peuvent être décimales, binaires (`b`), octales (`o`), hexadécimales (`h` ou préfixe `$`).
- Le compteur de position est `*`.

## 6. Directives prises en charge

| Directive | Rôle |
|---|---|
| `ORG` | Fixe l'adresse d'assemblage. |
| `END` | Termine un source ou un include. |
| `EQU` | Définit une constante symbolique. |
| `DB`, `DM` | Émettent des octets. |
| `DW` | Émet des mots 16 bits little-endian. |
| `DP` | Émet des adresses/pointeurs sur 3 octets. |
| `DS` | Réserve/remplit une zone. |
| `INCLUDE` | Inclut un fichier source. |
| `MACRO`, `ENDM` | Définissent une macro. |
| `LOCAL`, `ENDL` | Définissent une portée locale imbriquée. |
| `DEF`, `UNDEF` | Contrôlent les symboles conditionnels. |
| `IFDEF`, `IFNDEF`, `ELSE`, `ENDIF` | Assemblage conditionnel par symbole. |
| `IFEQ`, `IFNE`, `IFGT`, `IFLT` | Assemblage conditionnel par expression numérique. |
| `REPEAT`, `ENDR` | Répétition de blocs. |
| `STRUCT`, `ENDS` | Déclaration de structures et calcul de taille. |
| `SECTION` | Marque les sections pour MAP et rapport `-R`. |
| `PRE_ON`, `PRE_OFF` | Active/désactive les prébytes automatiques. |
| `PRE_PUSH`, `PRE_POP` | Sauvegarde/restaure l'état `PRE_ON/PRE_OFF`. |

## 7. Prébytes et nomenclature mémoire interne

La mémoire interne du SC62015 utilise des prébytes pour certaines combinaisons d'adressage. `xasm2026-4` respecte la table du manuel :

| 1er opérande \\ 2e opérande | `(n)` | `(BP+n)` | `(PY+n)` | `(BP+PY)` |
|---|---:|---:|---:|---:|
| `(n)` | `32h` | `30h` | `33h` | `31h` |
| `(BP+n)` | `22h` | interdit | `23h` | `21h` |
| `(PX+n)` | `36h` | `34h` | `37h` | `35h` |
| `(BP+PX)` | `26h` | `24h` | `27h` | `25h` |

Point essentiel confirmé pendant les validations : `PY+n` et `BP+PY` sont des formes de deuxième opérande, pas des formes de premier opérande. Les formes suivantes sont donc invalides et doivent produire `Prebyte error` :

```asm
mv (py+3),a
mv (bp+py),a
mv a,(py+3)
mv a,(bp+py)
```

Les formes suivantes sont valides parce que `PY` apparaît en deuxième opérande dans une instruction à deux opérandes internes :

```asm
mv  (bp+1),(py+2)     ; 23 C8 01 02
mvp (bp+px),(bp+py)   ; 25 CA 00 00
mvl (bp+1),(py+2)     ; 23 CB 01 02
```

## 8. Instructions SC62015 prises en charge

Les 69 mnémos documentés sont reconnus par le parser :

```text
NOP JP JPF JPZ JPNZ JPC JPNC JR JRZ JRNZ JRC JRNC CALL CALLF RET RETF RETI
PUSHU POPU PUSHS POPS
MV MVW MVP MVL MVLD EX EXW EXP EXL
ADD ADDB ADDW ADDP SUB SUBB SUBW SUBP ADC SBC ADCL SBCL DADL DSBL PMDF
CMP CMPW CMPP AND OR XOR TEST INC DEC ROR ROL SHR SHL DSRL DSLL SWAP
SC RC TCL HALT OFF IR WAIT RESET
```

Le coverage `tests/coverage_all.asm` vérifie plusieurs centaines de lignes de combinaisons d'instructions, directives, macros, conditions, sections, prébytes et sorties.

## 9. Sorties générées

| Sortie | Description |
|---|---|
| `.obj` | Objet XASM, avec types `Z`, `B`, `H`, `F` ou défaut. |
| `.lst` | Listing avec adresses, octets et source. En cas d'erreur fatale avec `-L`, reçoit un rapport d'erreur exploitable. |
| `.err` | Rapport d'erreur quand `-E` est actif. |
| `.hex` | Intel HEX, données utiles vérifiables par somme des records type 00. |
| `.s19` | Motorola S-Record. |
| `.map` | Sections et symboles. |
| `.d` | Dépendances, utile pour connaître les fichiers inclus. |
| `.uu` | BASIC auto-décodeur UUENCODE compatible Sharp, dérivé de `uuselfx.c`. |
| `.txt` | Dump texte type HxD. |

## 10. Rapport d'erreur

Depuis la finalisation du projet, une erreur fatale est écrite dans `.err` et `.lst` lorsque `-E` et `-L` sont utilisés. Exemple :

```text
XASM2026-4: ligne 3: Prebyte error | mv (py+3),a

bad_prebyte.asm    3    Prebyte error
    mv (py+3),a

Fatal error occured.
Assemble aborted.
```

Cela permet de corriger les erreurs sans perdre l'information de la console.

## 11. Travail réalisé sur xasm2026-4

Les principaux blocs finalisés sont :

- Port C# natif du parser de ligne de commande.
- Port des directives historiques et modernes.
- Port des macros, conditions, répétitions, includes et labels locaux imbriqués.
- Port des expressions numériques et du compteur `*`.
- Encodage complet des mnémos SC62015 documentés.
- Correction fine des prébytes, notamment la distinction entre colonnes `PY` et lignes `PX/BP+PX`.
- Alignement de `REGISTER`, `VOGUE`, `TRDOS`, `UUCODE` et `SAMPLES`.
- Générateur `.uu` aligné sur `uuselfx.c` (source de reference externe) : lignes BASIC, payload, checksum historique et ligne `size`.
- Rapport d'erreur écrit dans `.err` et `.lst`.
- Banc de couverture `coverage_all.asm`.

## 12. Validation et non-régression

Les validations effectuées en fin de projet :

| Groupe | Résultat |
|---|---|
| `SAMPLE1` à `SAMPLE4` | `.OBJ` identiques à la référence XASM 1.40. |
| `REGISTER` | `.OBJ` et `.LST` identiques à la référence. |
| `VOGUE` | `.OBJ` et `.LST` identiques à la référence. |
| `TRDOS` | `.OBJ` identiques au moteur C `xasm2026-2`. |
| `UUCODE` | `.OBJ` identiques au moteur C `xasm2026-2`. |
| `coverage_all.asm` | Compilation complète toutes options, 399 lignes, 892 octets. |
| HEX/S19 coverage | 892 octets utiles. |
| UU coverage | Ligne `size` identique à la taille réelle de l'objet. |

## 13. Commandes de validation recommandées

```powershell
cd tests
..\\bin\\xasm2026-4.exe coverage_all.asm -O coverage_all.obj -L coverage_all.lst -E -S -TZ -C -W -H -I coverage_all.hex -M coverage_all.s19 -P coverage_all.map -D coverage_all.d -B coverage_all.uu -X coverage_all.txt -V -R
```

Pour comparer les exemples, utiliser `fc.exe /b` sur les `.OBJ` et `fc.exe` sur les `.LST` quand une référence stricte existe.

## 14. Annexe A - Contenu complet de README - PC-E500 Instruction Table.md

{instruction}
"""


def build_xasm140_info_fr_markdown() -> str:
    original = read(XASM140_INFO_TXT).strip()
    today = date.today().strftime("%d/%m/%Y")
    return f"""# XASM140 - Information, version française enrichie xasm2026-4

Version du document : {today}

## 1. Introduction

XASM est un assembleur croisé absolu pour le CPU Sharp ESR-L / SC62015. Il a été créé par Narihito Kon et présenté dans l'univers des pocket computers Sharp, notamment la série PC-E500. Le CPU SC62015 est aussi appelé ESR-L. Des machines proches, comme certaines séries Organizer, Wizard ou Zaurus japonaises de génération compatible, exploitent des CPU apparentés.

La version historique XASM 1.40 a été améliorée à partir des versions 1.26 et 1.27j. Le moteur original, issu d'une conversion Turbo Pascal vers C, a notamment gagné en vitesse grâce à une gestion des symboles par hash.

`xasm2026-4` reprend cette base historique, mais remplace le coeur d'assemblage par un port C# natif maintenable, tout en conservant l'interface et les sorties utiles à la validation sur Sharp PC-E500S.

## 2. Caractéristiques de XASM 1.40

Les caractéristiques historiques restent au coeur du projet :

- assemblage absolu pour ESR-L / SC62015 ;
- génération automatique optionnelle des prébytes de RAM interne ;
- labels locaux imbriqués via `LOCAL` et `ENDL` ;
- fichiers inclus via `INCLUDE` ;
- macros avec `MACRO` et `ENDM` ;
- assemblage conditionnel via `IFDEF` et `IFNDEF` ;
- listing utilisable pour retrouver rapidement les erreurs ;
- plusieurs types de sortie objet ;
- compatibilité avec les workflows de transfert vers pocket computer.

## 3. Apports de xasm2026-4

`xasm2026-4` finalise un port C# natif en conservant le comportement utile de XASM 1.40 et de `xasm2026-2`.

| Domaine | Travail réalisé |
|---|---|
| Coeur assembleur | Port C# de `ORG`, `END`, `EQU`, données, macros, conditions, includes, sections, labels locaux et instructions. |
| Instructions | Les 69 mnémos documentés SC62015 sont reconnues et encodées. |
| Prébytes | Respect strict de la table `(n)`, `(BP+n)`, `(PX+n)`, `(BP+PX)` en premier opérande et `(n)`, `(BP+n)`, `(PY+n)`, `(BP+PY)` en second opérande. |
| UUENCODE | Générateur `-B` aligné sur `uuselfx.c`, y compris le payload BASIC, le checksum et la ligne `size`. |
| Diagnostics | Les erreurs fatales sont écrites dans `.err` et `.lst` quand `-E` et `-L` sont actifs. |
| Tests | `coverage_all.asm` couvre plusieurs centaines de lignes d'instructions et directives. |
| Régressions | `SAMPLES`, `REGISTER`, `VOGUE`, `TRDOS` et `UUCODE` validés. |

## 4. Syntaxe du source

Un source XASM contient une instruction ou directive par ligne. Le dernier énoncé du fichier principal doit être `END`. Les fichiers inclus historiques contiennent également leur propre `END`.

```asm
LABEL:  MNEMONIC  operand1,operand2   ; commentaire
        MNEMONIC  operand
        END
```

Les labels utilisent lettres, chiffres et `_`, ne commencent pas par un chiffre et sont insensibles à la casse. Les commentaires commencent par `;`.

## 5. Constantes, caractères et expressions

Les nombres sont évalués sur 20 bits. Sont acceptés :

| Notation | Base |
|---|---|
| `1010b` | binaire |
| `377o` | octal |
| `123d` ou `123` | décimal |
| `0BE000h` ou `$BE000` | hexadécimal |

Le caractère entre apostrophes vaut le code du dernier caractère :

```asm
'ABCD'  ; vaut le code de D
''''    ; apostrophe
''      ; code nul
```

Le compteur de position est `*`.

## 6. Directives d'assemblage

| Directive | Description |
|---|---|
| `ORG operand` | Fixe l'adresse courante. |
| `END` | Termine le fichier source ou inclus. |
| `LABEL: EQU operand` | Définit un symbole. |
| `DB`, `DM`, `DW`, `DP`, `DS` | Émettent ou réservent des données. |
| `PRE operand` | Émet un prébyte manuel. |
| `PRE_ON`, `PRE_OFF` | Active ou désactive la génération automatique de prébytes. |
| `PRE_PUSH`, `PRE_POP` | Sauvegarde/restaure l'état de génération automatique. |
| `INCLUDE` | Inclut un fichier. |
| `MACRO`, `ENDM` | Définit une macro. |
| `LOCAL`, `ENDL` | Ouvre/ferme une portée locale. |
| `DEF`, `UNDEF`, `IFDEF`, `IFNDEF`, `ELSE`, `ENDIF` | Assemblage conditionnel historique. |
| `IFEQ`, `IFNE`, `IFGT`, `IFLT` | Assemblage conditionnel numérique ajouté dans la lignée 2026. |
| `REPEAT`, `ENDR` | Répète un bloc. |
| `STRUCT`, `ENDS` | Décrit une structure et calcule sa taille. |
| `SECTION` | Marque une section pour MAP et rapport de taille. |

## 7. Prébytes automatiques

La génération automatique est effective avec `PRE_ON`. Le mode historique par défaut est `PRE_OFF`.

La table valide est :

| 1er opérande \\ 2e opérande | `(n)` | `(BP+n)` | `(PY+n)` | `(BP+PY)` |
|---|---:|---:|---:|---:|
| `(n)` | `32h` | `30h` | `33h` | `31h` |
| `(BP+n)` | `22h` | interdit | `23h` | `21h` |
| `(PX+n)` | `36h` | `34h` | `37h` | `35h` |
| `(BP+PX)` | `26h` | `24h` | `27h` | `25h` |

Attention : `PY+n` et `BP+PY` ne sont pas des formes de premier opérande. Elles ne doivent pas être utilisées comme destination ou opérande unique. `xasm2026-4` renvoie désormais `Prebyte error` pour ces cas.

## 8. Options de xasm2026-4

```text
-L  listing
-E  rapport d'erreur
-O  objet
-S  symboles dans le listing
-T  type objet : Z, B, H, F ou défaut
-C  compteur de lignes
-W  avertissements
-H  hash off historique
-I  Intel HEX
-M  Motorola S-Record
-P  MAP
-D  dépendances
-B  BASIC UUENCODE auto-décodeur
-X  dump texte HxD
-V  erreurs détaillées
-R  rapport de sections
```

## 9. Diagnostics et correction des erreurs

Quand `-L -E` sont utilisés et qu'une erreur fatale survient, `xasm2026-4` écrit un rapport dans `.err` et `.lst`.

Exemple :

```text
XASM2026-4: ligne 3: Prebyte error | mv (py+3),a

bad_prebyte.asm    3    Prebyte error
    mv (py+3),a

Fatal error occured.
Assemble aborted.
```

Cela facilite la correction séquentielle des erreurs source.

## 10. Validation finale

| Ensemble | Résultat |
|---|---|
| `SAMPLE1` à `SAMPLE4` | `.OBJ` identiques à la référence. |
| `REGISTER` | `.OBJ/.LST` identiques à XASM 1.40. |
| `VOGUE` | `.OBJ/.LST` identiques à XASM 1.40. |
| `TRDOS` | `.OBJ` identiques au moteur C. |
| `UUCODE` | `.OBJ` identiques au moteur C. |
| `coverage_all.asm` | Compilation toutes options, 399 lignes, 892 octets. |

## 11. Texte original anglais de XASM140 - Information.DOC

Le texte ci-dessous est conservé comme référence historique.

{fenced(original)}
"""


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def format_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [
        ("Heading 1", 16, RGBColor(31, 77, 120)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 11, RGBColor(31, 77, 120)),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def add_code_paragraph(doc: Document, line: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(line if line else " ")
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(normalized_rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value.strip()
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if row_index == 0:
                set_cell_shading(cells[index], "E8EEF5")
                for paragraph in cells[index].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", cell.strip()) for cell in row):
            rows.append(row)
        i += 1
    return rows, i


def markdown_to_docx(markdown: str, out_path: Path, title: str) -> None:
    doc = Document()
    format_document(doc)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(11, 37, 69)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Document généré le {date.today().strftime('%d/%m/%Y')}").italic = True
    doc.add_paragraph()

    lines = markdown.splitlines()
    in_code = False
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_code_paragraph(doc, line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(stripped[2:].strip())
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            run.font.color.rgb = RGBColor(90, 90, 90)
        else:
            doc.add_paragraph(stripped)
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> None:
    """Produit les .docx a partir des .md de Documentation/.

    Les fichiers Markdown sont **maintenus a la main** et font foi : ils sont versionnes et
    relus. Le generateur ne fait donc, par defaut, que produire les .docx correspondants.

    Auparavant il reconstruisait aussi les .md depuis un gabarit code en dur dans ce
    fichier, ce qui ecrasait silencieusement toute modification apportee a la
    documentation. Ce comportement reste accessible via --rebuild-markdown, pour le cas ou
    l'on voudrait repartir du gabarit, mais il n'est plus le chemin par defaut.
    """
    rebuild_markdown = "--rebuild-markdown" in sys.argv

    DOC_DIR.mkdir(parents=True, exist_ok=True)
    xasm_doc_md = DOC_DIR / "Documentation_XASM2026-4_PC-E500S.md"
    xasm_doc_docx = DOC_DIR / "Documentation_XASM2026-4_PC-E500S.docx"
    info_md = DOC_DIR / "XASM140_Information_FR_xasm2026-4.md"
    info_docx = DOC_DIR / "XASM140_Information_FR_xasm2026-4.docx"

    if rebuild_markdown:
        print("!! --rebuild-markdown : les .md vont etre ecrases par le gabarit interne.")
        xasm_markdown = build_xasm2026_3_markdown()
        info_markdown = build_xasm140_info_fr_markdown()
        write(xasm_doc_md, xasm_markdown)
        write(info_md, info_markdown)
    else:
        for source in (xasm_doc_md, info_md):
            if not source.exists():
                raise SystemExit(
                    f"Markdown source introuvable : {source}\n"
                    "Utiliser --rebuild-markdown pour le reconstruire depuis le gabarit."
                )

        xasm_markdown = read(xasm_doc_md)
        info_markdown = read(info_md)

    markdown_to_docx(xasm_markdown, xasm_doc_docx, "Documentation XASM2026-4 pour Sharp PC-E500S")
    markdown_to_docx(info_markdown, info_docx, "XASM140 - Information, version française enrichie")

    FINAL_DOC_DIR.mkdir(parents=True, exist_ok=True)
    for path in [xasm_doc_md, xasm_doc_docx, info_md, info_docx]:
        target = FINAL_DOC_DIR / path.name
        target.write_bytes(path.read_bytes())

    print(xasm_doc_md)
    print(xasm_doc_docx)
    print(info_md)
    print(info_docx)


if __name__ == "__main__":
    main()
